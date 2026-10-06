import json
import subprocess
from collections.abc import Iterable
from dataclasses import asdict
from datetime import datetime, timezone
from importlib import metadata
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from spectral_ibr.application.dto.ibr_case_build import (
    IbrBuildRequest,
    IbrBuildResult,
    IbrConversionProtocol,
    IbrReplacementGroup,
)
from spectral_ibr.domain.errors import (
    InvalidConversionProfileError,
    IbrConversionError,
    MissingBaseCaseError,
)

REGCP1_HEADERS = [
    "uid",
    "idx",
    "u",
    "name",
    "bus",
    "gen",
    "Sn",
    "Tg",
    "Rrpwr",
    "Brkpt",
    "Zerox",
    "Lvplsw",
    "Lvpl1",
    "Volim",
    "Lvpnt1",
    "Lvpnt0",
    "Iolim",
    "Tfltr",
    "Khv",
    "Iqrmax",
    "Iqrmin",
    "Accel",
    "gammap",
    "gammaq",
    "pll",
]
REECA1_HEADERS = [
    "uid",
    "idx",
    "u",
    "name",
    "reg",
    "busr",
    "PFFLAG",
    "VFLAG",
    "QFLAG",
    "PFLAG",
    "PQFLAG",
    "Vdip",
    "Vup",
    "Trv",
    "dbd1",
    "dbd2",
    "Kqv",
    "Iqh1",
    "Iql1",
    "Vref0",
    "Iqfrz",
    "Thld",
    "Thld2",
    "Tp",
    "QMax",
    "QMin",
    "VMAX",
    "VMIN",
    "Kqp",
    "Kqi",
    "Kvp",
    "Kvi",
    "Vref1",
    "Tiq",
    "dPmax",
    "dPmin",
    "PMAX",
    "PMIN",
    "Imax",
    "Tpord",
    "Vq1",
    "Iq1",
    "Vq2",
    "Iq2",
    "Vq3",
    "Iq3",
    "Vq4",
    "Iq4",
    "Vp1",
    "Ip1",
    "Vp2",
    "Ip2",
    "Vp3",
    "Ip3",
    "Vp4",
    "Ip4",
]
REPCA1_HEADERS = [
    "uid",
    "idx",
    "u",
    "name",
    "ree",
    "line",
    "busr",
    "busf",
    "VCFlag",
    "RefFlag",
    "Fflag",
    "PLflag",
    "Tfltr",
    "Kp",
    "Ki",
    "Tft",
    "Tfv",
    "Vfrz",
    "Rc",
    "Xc",
    "Kc",
    "emax",
    "emin",
    "dbd1",
    "dbd2",
    "Qmax",
    "Qmin",
    "Kpg",
    "Kig",
    "Tp",
    "fdbd1",
    "fdbd2",
    "femax",
    "femin",
    "Pmax",
    "Pmin",
    "Tg",
    "Ddn",
    "Dup",
]
PLL1_HEADERS = ["uid", "idx", "u", "name", "bus", "Kp", "Ki", "Tf", "Tp", "fn"]
REGF1_HEADERS = [
    "uid",
    "idx",
    "u",
    "name",
    "bus",
    "gen",
    "Sn",
    "rf",
    "xf",
    "Vdip",
    "Tfrz",
    "PQFLAG",
    "fn",
    "dwmax",
    "dwmin",
    "wdrp",
    "Qdrp",
    "Tr",
    "Te",
    "KPi",
    "KIi",
    "KPv",
    "KIv",
    "Pmax",
    "Pmin",
    "KPplim",
    "KIplim",
    "Qmax",
    "Qmin",
    "KPqlim",
    "KIqlim",
    "Tpm",
    "gammap",
    "gammaq",
]


class SgToIbrReplacer:
    """Deterministic ANDES SG-to-IBR conversion adapter."""

    def __init__(self, project_root: Path = Path.cwd()) -> None:
        self.project_root = project_root

    def build(
        self,
        request: IbrBuildRequest,
        protocol: IbrConversionProtocol,
    ) -> IbrBuildResult:
        output_dir = request.output_dir
        output_dir.mkdir(parents=True, exist_ok=True)

        resolved_base_case = self._resolve_base_case(protocol)
        processed_case = self._resolve_project_path(protocol.output_case)
        processed_case.parent.mkdir(parents=True, exist_ok=True)

        workbook = load_workbook(resolved_base_case)
        conversion_plan = self._convert_workbook(workbook, protocol)
        workbook.save(processed_case)

        plan_path = output_dir / "conversion_plan.json"
        protocol_snapshot_path = output_dir / "conversion_protocol_snapshot.json"
        manifest_path = output_dir / "manifest.json"

        plan = self._build_plan(protocol, resolved_base_case, processed_case, conversion_plan)
        protocol_snapshot = self._protocol_to_json(protocol)
        manifest = self._build_manifest(
            request,
            protocol,
            resolved_base_case,
            processed_case,
            plan_path,
        )

        plan_path.write_text(json.dumps(plan, indent=2), encoding="utf-8")
        protocol_snapshot_path.write_text(
            json.dumps(protocol_snapshot, indent=2),
            encoding="utf-8",
        )
        manifest_path.write_text(json.dumps(manifest, indent=2), encoding="utf-8")

        return IbrBuildResult(
            case_name=request.case_name,
            status="built",
            output_dir=output_dir,
            manifest_path=manifest_path,
            artifacts={
                "processed_case": processed_case,
                "conversion_plan": plan_path,
                "protocol_snapshot": protocol_snapshot_path,
                "manifest": manifest_path,
            },
            metrics={
                "replacement_groups": float(len(protocol.replacements)),
                "converted_generators": float(
                    sum(len(group.gens) for group in protocol.replacements)
                ),
                "eligible_buses": float(len(protocol.eligible_buses)),
            },
            message=f"Generated processed ANDES workbook: {processed_case}",
        )

    def _resolve_base_case(self, protocol: IbrConversionProtocol) -> Path:
        direct = self._resolve_project_path(protocol.base_case)
        if direct.exists():
            return direct
        raise MissingBaseCaseError(
            f"Configured base_case does not exist: {direct}. "
            "No fallback is used; put the source workbook there or update config."
        )

    def _resolve_project_path(self, path: Path) -> Path:
        if path.is_absolute():
            return path
        return self.project_root / path

    def _convert_workbook(
        self,
        workbook: Any,
        protocol: IbrConversionProtocol,
    ) -> dict[str, Any]:
        converted_pairs = self._converted_pairs(protocol.replacements)
        converted_gens = {gen for gen, _ in converted_pairs}
        static_generators = self._static_generators(workbook)
        self._validate_replacements(protocol, converted_pairs, static_generators)

        removed = self._remove_replaced_synchronous_dynamics(workbook, converted_gens)
        busfreq_by_bus = self._busfreq_by_bus(workbook)

        added: dict[str, list[str]] = {
            "REGCP1": [],
            "REECA1": [],
            "REPCA1": [],
            "PLL1": [],
            "REGF1": [],
        }
        for group in protocol.replacements:
            for gen, bus in zip(group.gens, group.buses, strict=True):
                static = static_generators[gen]
                model = group.model.upper()
                if model == "REGCP1":
                    ids = self._append_gfl_chain(workbook, protocol, group, gen, bus, static, busfreq_by_bus)
                    for sheet_name, idx in ids.items():
                        added[sheet_name].append(idx)
                elif model == "REGF1":
                    idx = self._append_regf1(workbook, protocol, group, gen, bus, static)
                    added["REGF1"].append(idx)
                else:
                    raise IbrConversionError(
                        f"Unsupported IBR model '{group.model}' in profile {protocol.profile_name}."
                    )

        return {
            "profile": protocol.profile_name,
            "converted_pairs": [{"gen": gen, "bus": bus} for gen, bus in converted_pairs],
            "removed": removed,
            "added": added,
        }

    def _converted_pairs(
        self,
        replacements: Iterable[IbrReplacementGroup],
    ) -> list[tuple[int, int]]:
        pairs: list[tuple[int, int]] = []
        seen: set[int] = set()
        for group in replacements:
            if len(group.gens) != len(group.buses):
                raise InvalidConversionProfileError(
                    f"Replacement group '{group.kind}' has mismatched gens and buses."
                )
            for gen, bus in zip(group.gens, group.buses, strict=True):
                if gen in seen:
                    raise InvalidConversionProfileError(
                        f"Generator {gen} is assigned more than once."
                    )
                seen.add(gen)
                pairs.append((gen, bus))
        return pairs

    def _static_generators(self, workbook: Any) -> dict[int, dict[str, Any]]:
        static: dict[int, dict[str, Any]] = {}
        for sheet_name in ("PV", "Slack"):
            if sheet_name not in workbook.sheetnames:
                continue
            for record in self._records(workbook[sheet_name]):
                idx = record.get("idx")
                if idx is not None:
                    static[int(idx)] = record
        return static

    def _validate_replacements(
        self,
        protocol: IbrConversionProtocol,
        converted_pairs: list[tuple[int, int]],
        static_generators: dict[int, dict[str, Any]],
    ) -> None:
        eligible = set(protocol.eligible_buses)
        keep = set(protocol.keep_synchronous)
        for gen, bus in converted_pairs:
            if gen not in static_generators:
                raise InvalidConversionProfileError(
                    f"Generator {gen} is not present in PV/Slack sheets."
                )
            actual_bus = int(static_generators[gen]["bus"])
            if actual_bus != bus:
                raise InvalidConversionProfileError(
                    f"Generator {gen} is configured at bus {bus}, "
                    f"but the base case has it at bus {actual_bus}."
                )
            if bus not in eligible:
                raise InvalidConversionProfileError(
                    f"Bus {bus} is not in eligible_buses for protocol {protocol.name}."
                )
            if bus in keep:
                raise InvalidConversionProfileError(
                    f"Bus {bus} is marked keep_synchronous and cannot be converted."
                )

    def _remove_replaced_synchronous_dynamics(
        self,
        workbook: Any,
        converted_gens: set[int],
    ) -> dict[str, list[str]]:
        removed: dict[str, list[str]] = {}
        removed_syn = self._filter_sheet(
            workbook,
            "GENROU",
            lambda row: int(row.get("gen")) in converted_gens,
        )
        removed["GENROU"] = [str(row["idx"]) for row in removed_syn]
        removed_syn_idx = set(removed["GENROU"])

        for sheet_name in ("TGOV1N", "TGOV1"):
            removed_rows = self._filter_sheet(
                workbook,
                sheet_name,
                lambda row: str(row.get("syn")) in removed_syn_idx,
            )
            if removed_rows:
                removed[sheet_name] = [str(row["idx"]) for row in removed_rows]

        removed_avr = self._filter_sheet(
            workbook,
            "IEEEX1",
            lambda row: str(row.get("syn")) in removed_syn_idx,
        )
        removed["IEEEX1"] = [str(row["idx"]) for row in removed_avr]
        removed_avr_idx = set(removed["IEEEX1"])

        removed_pss = self._filter_sheet(
            workbook,
            "IEEEST",
            lambda row: str(row.get("avr")) in removed_avr_idx,
        )
        removed["IEEEST"] = [str(row["idx"]) for row in removed_pss]
        return removed

    def _filter_sheet(
        self,
        workbook: Any,
        sheet_name: str,
        predicate: Any,
    ) -> list[dict[str, Any]]:
        if sheet_name not in workbook.sheetnames:
            return []
        worksheet = workbook[sheet_name]
        headers = self._headers(worksheet)
        records = self._records(worksheet)
        removed = [record for record in records if predicate(record)]
        kept = [record for record in records if not predicate(record)]
        self._write_records(worksheet, headers, kept)
        return removed

    def _append_gfl_chain(
        self,
        workbook: Any,
        protocol: IbrConversionProtocol,
        group: IbrReplacementGroup,
        gen: int,
        bus: int,
        static: dict[str, Any],
        busfreq_by_bus: dict[int, str],
    ) -> dict[str, str]:
        reg_idx = f"REGCP1_G{gen}"
        ree_idx = f"REECA1_G{gen}"
        rep_idx = f"REPCA1_G{gen}"
        pll_idx = f"PLL1_G{gen}" if group.with_pll else None
        parameter_set = group.parameter_set

        if pll_idx is not None:
            self._append_record(
                workbook,
                "PLL1",
                PLL1_HEADERS,
                self._with_overrides(
                    protocol,
                    parameter_set,
                    "PLL1",
                    {
                        "idx": pll_idx,
                        "name": pll_idx,
                        "u": 1,
                        "bus": bus,
                        "Kp": 1,
                        "Ki": 0.2,
                        "Tf": 0.05,
                        "Tp": 0.05,
                        "fn": 60,
                    },
                ),
            )

        self._append_record(
            workbook,
            "REGCP1",
            REGCP1_HEADERS,
            self._with_overrides(
                protocol,
                parameter_set,
                "REGCP1",
                {
                    "idx": reg_idx,
                    "name": reg_idx,
                    "u": 1,
                    "bus": bus,
                    "gen": gen,
                    "Sn": static.get("Sn"),
                    "Tg": 0.1,
                    "Rrpwr": 10,
                    "Brkpt": 1,
                    "Zerox": 0.5,
                    "Lvplsw": 1,
                    "Lvpl1": 1,
                    "Volim": 1.2,
                    "Lvpnt1": 0.8,
                    "Lvpnt0": 0.4,
                    "Iolim": -1.5,
                    "Tfltr": 0.1,
                    "Khv": 0.7,
                    "Iqrmax": 1,
                    "Iqrmin": -1,
                    "Accel": 0,
                    "gammap": 1,
                    "gammaq": 1,
                    "pll": pll_idx,
                },
            ),
        )
        self._append_record(
            workbook,
            "REECA1",
            REECA1_HEADERS,
            self._with_overrides(
                protocol,
                parameter_set,
                "REECA1",
                {
                    "idx": ree_idx,
                    "name": ree_idx,
                    "u": 1,
                    "reg": reg_idx,
                    "busr": None,
                    "PFFLAG": 0,
                    "VFLAG": 0,
                    "QFLAG": 0,
                    "PFLAG": 0,
                    "PQFLAG": 0,
                    "Vdip": 0.8,
                    "Vup": 1.1,
                    "Trv": 0.02,
                    "dbd1": -0.02,
                    "dbd2": 0.02,
                    "Kqv": 20,
                    "Iqh1": 999,
                    "Iql1": -999,
                    "Vref0": static.get("v0", 1),
                    "Iqfrz": 0,
                    "Thld": -2,
                    "Thld2": 1,
                    "Tp": 0.02,
                    "QMax": 999,
                    "QMin": -999,
                    "VMAX": 999,
                    "VMIN": -999,
                    "Kqp": 1,
                    "Kqi": 0.1,
                    "Kvp": 1,
                    "Kvi": 0.1,
                    "Vref1": static.get("v0", 1),
                    "Tiq": 0.02,
                    "dPmax": 999,
                    "dPmin": -999,
                    "PMAX": 999,
                    "PMIN": -999,
                    "Imax": 10,
                    "Tpord": 0.02,
                    "Vq1": 0.2,
                    "Iq1": 2,
                    "Vq2": 0.4,
                    "Iq2": 4,
                    "Vq3": 0.8,
                    "Iq3": 8,
                    "Vq4": 1,
                    "Iq4": 10,
                    "Vp1": 0.2,
                    "Ip1": 2,
                    "Vp2": 0.4,
                    "Ip2": 4,
                    "Vp3": 0.8,
                    "Ip3": 8,
                    "Vp4": 1,
                    "Ip4": 10,
                },
            ),
        )
        self._append_record(
            workbook,
            "REPCA1",
            REPCA1_HEADERS,
            self._with_overrides(
                protocol,
                parameter_set,
                "REPCA1",
                {
                    "idx": rep_idx,
                    "name": rep_idx,
                    "u": 1,
                    "ree": ree_idx,
                    "line": None,
                    "busr": None,
                    "busf": busfreq_by_bus.get(bus),
                    "VCFlag": 1,
                    "RefFlag": 1,
                    "Fflag": 0,
                    "PLflag": 0,
                    "Tfltr": 0.02,
                    "Kp": 1,
                    "Ki": 0.1,
                    "Tft": 1,
                    "Tfv": 1,
                    "Vfrz": 0.8,
                    "Rc": None,
                    "Xc": None,
                    "Kc": 1,
                    "emax": 999,
                    "emin": -999,
                    "dbd1": -0.02,
                    "dbd2": 0.02,
                    "Qmax": 999,
                    "Qmin": -999,
                    "Kpg": 1,
                    "Kig": 0.1,
                    "Tp": 0.02,
                    "fdbd1": -0.01,
                    "fdbd2": 0.01,
                    "femax": 0.05,
                    "femin": -0.05,
                    "Pmax": 999,
                    "Pmin": -999,
                    "Tg": 0.02,
                    "Ddn": 10,
                    "Dup": 10,
                },
            ),
        )
        ids = {"REGCP1": reg_idx, "REECA1": ree_idx, "REPCA1": rep_idx}
        if pll_idx is not None:
            ids["PLL1"] = pll_idx
        return ids

    def _append_regf1(
        self,
        workbook: Any,
        protocol: IbrConversionProtocol,
        group: IbrReplacementGroup,
        gen: int,
        bus: int,
        static: dict[str, Any],
    ) -> str:
        idx = f"REGF1_G{gen}"
        self._append_record(
            workbook,
            "REGF1",
            REGF1_HEADERS,
            self._with_overrides(
                protocol,
                group.parameter_set,
                "REGF1",
                {
                    "idx": idx,
                    "name": idx,
                    "u": 1,
                    "bus": bus,
                    "gen": gen,
                    "Sn": static.get("Sn"),
                    "rf": static.get("ra", 0),
                    "xf": static.get("xs", 0.2),
                    "Vdip": 0.8,
                    "Tfrz": 0,
                    "PQFLAG": 1,
                    "fn": 60,
                    "dwmax": 75,
                    "dwmin": -75,
                    "wdrp": 0.033,
                    "Qdrp": 0.045,
                    "Tr": 0.005,
                    "Te": 0.005,
                    "KPi": 0.5,
                    "KIi": 20,
                    "KPv": 3,
                    "KIv": 10,
                    "Pmax": static.get("pmax", 1),
                    "Pmin": static.get("pmin", -1),
                    "KPplim": 5,
                    "KIplim": 30,
                    "Qmax": static.get("qmax", 1),
                    "Qmin": static.get("qmin", -1),
                    "KPqlim": 0.1,
                    "KIqlim": 1.5,
                    "Tpm": 0.025,
                    "gammap": 1,
                    "gammaq": 1,
                },
            ),
        )
        return idx

    def _busfreq_by_bus(self, workbook: Any) -> dict[int, str]:
        if "BusFreq" not in workbook.sheetnames:
            return {}
        return {
            int(record["bus"]): str(record["idx"])
            for record in self._records(workbook["BusFreq"])
            if record.get("bus") is not None and record.get("idx") is not None
        }

    def _append_record(
        self,
        workbook: Any,
        sheet_name: str,
        headers: list[str],
        record: dict[str, Any],
    ) -> None:
        worksheet = self._ensure_sheet(workbook, sheet_name, headers)
        current_headers = self._headers(worksheet)
        row = worksheet.max_row + 1
        record = {"uid": self._next_uid(worksheet), **record}
        for column, header in enumerate(current_headers, start=1):
            worksheet.cell(row=row, column=column, value=record.get(header))

    def _ensure_sheet(self, workbook: Any, sheet_name: str, headers: list[str]) -> Worksheet:
        if sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            if worksheet.max_row == 0 or not self._headers(worksheet):
                for column, header in enumerate(headers, start=1):
                    worksheet.cell(row=1, column=column, value=header)
            return worksheet
        worksheet = workbook.create_sheet(sheet_name)
        for column, header in enumerate(headers, start=1):
            worksheet.cell(row=1, column=column, value=header)
        return worksheet

    def _with_overrides(
        self,
        protocol: IbrConversionProtocol,
        parameter_set: str | None,
        model: str,
        defaults: dict[str, Any],
    ) -> dict[str, Any]:
        if parameter_set is None:
            return defaults
        models = protocol.parameter_sets.get(parameter_set, {})
        overrides = models.get(model, models.get(model.lower(), {}))
        return {**defaults, **overrides}

    def _headers(self, worksheet: Worksheet) -> list[str]:
        return [
            str(worksheet.cell(row=1, column=column).value)
            for column in range(1, worksheet.max_column + 1)
            if worksheet.cell(row=1, column=column).value is not None
        ]

    def _records(self, worksheet: Worksheet) -> list[dict[str, Any]]:
        headers = self._headers(worksheet)
        records: list[dict[str, Any]] = []
        for row in range(2, worksheet.max_row + 1):
            values = [worksheet.cell(row=row, column=column).value for column in range(1, len(headers) + 1)]
            if not any(value is not None for value in values):
                continue
            records.append(dict(zip(headers, values, strict=True)))
        return records

    def _write_records(
        self,
        worksheet: Worksheet,
        headers: list[str],
        records: list[dict[str, Any]],
    ) -> None:
        if worksheet.max_row > 1:
            worksheet.delete_rows(2, worksheet.max_row - 1)
        for row_number, record in enumerate(records, start=2):
            for column, header in enumerate(headers, start=1):
                worksheet.cell(row=row_number, column=column, value=record.get(header))

    def _next_uid(self, worksheet: Worksheet) -> int:
        headers = self._headers(worksheet)
        if "uid" not in headers:
            return 0
        records = self._records(worksheet)
        uid_values = [record.get("uid") for record in records if record.get("uid") is not None]
        int_values = []
        for value in uid_values:
            try:
                int_values.append(int(value))
            except (TypeError, ValueError):
                continue
        return max(int_values, default=-1) + 1

    def _build_plan(
        self,
        protocol: IbrConversionProtocol,
        resolved_base_case: Path,
        processed_case: Path,
        conversion_plan: dict[str, Any],
    ) -> dict[str, Any]:
        return {
            "protocol": protocol.name,
            "profile": protocol.profile_name,
            "base_case": str(resolved_base_case),
            "output_case": str(processed_case),
            "preserve_power_flow": protocol.preserve_power_flow,
            "eligible_buses": list(protocol.eligible_buses),
            "keep_synchronous": list(protocol.keep_synchronous),
            "remove_dynamic_models": ["GENROU", "governor", "exciter", "PSS"],
            "add_dynamic_models": sorted({group.model for group in protocol.replacements}),
            "replacements": [asdict(group) for group in protocol.replacements],
            "state_partition_policy": {
                "electrical": "network/interface states",
                "control": "states matching PLL/REGC/REGCP/REGF/REEC/REPC/IBR/WT",
            },
            "conversion": conversion_plan,
            "status": "built",
        }

    def _build_manifest(
        self,
        request: IbrBuildRequest,
        protocol: IbrConversionProtocol,
        resolved_base_case: Path,
        processed_case: Path,
        plan_path: Path,
    ) -> dict[str, Any]:
        return {
            "case": request.case_name,
            "profile": request.profile_name,
            "protocol": protocol.name,
            "created_at_utc": datetime.now(timezone.utc).isoformat(),
            "seed": request.seed,
            "git_commit": self._git_commit(),
            "andes_version": self._andes_version(),
            "base_case": str(resolved_base_case),
            "processed_case": str(processed_case),
            "protocol_path": str(request.protocol_path),
            "plan_path": str(plan_path),
            "citations": list(protocol.citations),
            "status": "built",
        }

    def _protocol_to_json(self, protocol: IbrConversionProtocol) -> dict[str, Any]:
        payload = asdict(protocol)
        payload["base_case"] = str(protocol.base_case)
        payload["output_case"] = str(protocol.output_case)
        return payload

    def _git_commit(self) -> str | None:
        try:
            result = subprocess.run(
                ["git", "rev-parse", "HEAD"],
                cwd=self.project_root,
                check=True,
                capture_output=True,
                text=True,
            )
        except (FileNotFoundError, subprocess.CalledProcessError):
            return None
        return result.stdout.strip()

    def _andes_version(self) -> str | None:
        try:
            return metadata.version("andes")
        except metadata.PackageNotFoundError:
            return None


def protocol_from_mapping(
    mapping: dict[str, Any],
    profile_name: str | None = None,
) -> IbrConversionProtocol:
    selected_profile = profile_name or mapping.get("active_profile")
    profiles = mapping.get("profiles", {})
    if not selected_profile:
        raise InvalidConversionProfileError(
            "No profile selected. Set active_profile in config or pass --profile."
        )
    if selected_profile not in profiles:
        available = ", ".join(sorted(profiles)) or "<none>"
        raise InvalidConversionProfileError(
            f"Unknown conversion profile '{selected_profile}'. Available: {available}."
        )

    profile = profiles[selected_profile]
    replacements = _replacement_groups(profile.get("replacements", {}))
    base_case = mapping.get("base_case")
    if isinstance(base_case, dict):
        base_case_path = Path(base_case["path"])
    else:
        base_case_path = Path(base_case)

    return IbrConversionProtocol(
        name=str(mapping["name"]),
        base_case=base_case_path,
        output_case=Path(profile["output_case"]),
        profile_name=str(selected_profile),
        eligible_buses=tuple(mapping.get("eligible_buses", ())),
        keep_synchronous=tuple(profile.get("keep_synchronous", ())),
        replacements=tuple(replacements),
        preserve_power_flow=bool(mapping.get("preserve_power_flow", True)),
        parameter_sets=mapping.get("parameter_sets", {}),
        citations=tuple(mapping.get("citations", ())),
        notes=tuple(mapping.get("notes", ())),
    )


def _replacement_groups(raw_replacements: dict[str, dict[str, Any]]) -> list[IbrReplacementGroup]:
    replacements = []
    for kind, group in raw_replacements.items():
        replacements.append(
            IbrReplacementGroup(
                kind=kind,
                gens=tuple(group.get("gens", ())),
                buses=tuple(group.get("buses", ())),
                model=str(group["model"]).upper(),
                with_pll=bool(group.get("with_pll", False)),
                parameter_set=group.get("parameter_set"),
                source=group.get("source"),
            )
        )
    return replacements
