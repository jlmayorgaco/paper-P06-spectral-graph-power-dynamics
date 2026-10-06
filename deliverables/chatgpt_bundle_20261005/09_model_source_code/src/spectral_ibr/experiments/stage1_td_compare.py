import json
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import matplotlib.pyplot as plt
import pandas as pd
import yaml
from openpyxl import load_workbook

from spectral_ibr.domain.errors import MissingBaseCaseError
from spectral_ibr.experiments.base_gate import BaseGate, GateResult, GateRunConfig
from spectral_ibr.infrastructure.andes.sg_to_ibr_replacer import protocol_from_mapping


@dataclass(frozen=True)
class SimulationSeries:
    name: str
    frequency: pd.DataFrame
    voltage: pd.DataFrame
    gen_omega: pd.DataFrame
    regcp1_pe: pd.DataFrame | None
    regf1_pe: pd.DataFrame | None
    final_time: float


@dataclass(frozen=True)
class LoadToggleEvent:
    load: str = "PQ_7"
    time: float = 12.0
    final_time: float = 20.0

    @property
    def label(self) -> str:
        return f"PQ load {self.load} disconnect at t={self.time:g} s"


class Stage1TdCompareGate(BaseGate):
    name = "stage1-td-compare"

    def __init__(self, project_root: Path | None = None) -> None:
        self.project_root = project_root or Path.cwd()

    def run(self, config: GateRunConfig | None = None) -> GateResult:
        config = config or GateRunConfig(name=self.name, profile_name="mix60")
        protocol_path = self._resolve(
            config.config_path or Path("configs/cases/conversion_protocol.yaml")
        )
        mapping = yaml.safe_load(protocol_path.read_text(encoding="utf-8"))
        protocol = protocol_from_mapping(mapping, profile_name=config.profile_name)
        profile = protocol.profile_name

        base_case = self._resolve(protocol.base_case)
        processed_case = self._resolve(protocol.output_case)
        if not base_case.exists():
            raise MissingBaseCaseError(
                f"Configured base_case does not exist: {base_case}."
            )
        if not processed_case.exists():
            raise MissingBaseCaseError(
                f"Processed case does not exist: {processed_case}. "
                f"Run stage1-build-ibr-case --profile {profile} first."
            )

        output_dir = config.output_root / "stage1_td_compare" / profile
        output_dir.mkdir(parents=True, exist_ok=True)
        event = LoadToggleEvent()

        base_scenario = self._write_load_toggle_scenario(
            base_case,
            output_dir / "base_toggle_pq7.xlsx",
            event,
        )
        ibr_scenario = self._write_load_toggle_scenario(
            processed_case,
            output_dir / f"{profile}_toggle_pq7.xlsx",
            event,
        )

        base = self._simulate("base", base_scenario, event.final_time)
        ibr = self._simulate(profile, ibr_scenario, event.final_time)
        artifacts = self._write_outputs(output_dir, base, ibr, profile, event)
        artifacts["base_scenario_case"] = base_scenario
        artifacts["ibr_scenario_case"] = ibr_scenario
        metrics = self._metrics(base, ibr)
        completed = self._completed(base, ibr, event.final_time)
        verdict = "passed" if completed else "failed"
        message = (
            f"Dynamic comparison written to {output_dir}"
            if completed
            else (
                "Dynamic comparison stopped early. "
                f"base final t={base.final_time:.4f} s, "
                f"{profile} final t={ibr.final_time:.4f} s. See {output_dir}."
            )
        )

        return GateResult(
            name=self.name,
            verdict=verdict,
            artifacts={key: str(path) for key, path in artifacts.items()},
            metrics=metrics,
            message=message,
        )

    def _write_load_toggle_scenario(
        self,
        source_case: Path,
        destination: Path,
        event: LoadToggleEvent,
    ) -> Path:
        workbook = load_workbook(source_case)
        self._disable_togglers(workbook)
        if "Alter" in workbook.sheetnames:
            del workbook["Alter"]
        if "Toggler" in workbook.sheetnames:
            sheet = workbook["Toggler"]
        else:
            sheet = workbook.create_sheet("Toggler")
            sheet.append(["uid", "idx", "u", "name", "model", "dev", "t"])
        sheet.append(
            [
                sheet.max_row - 1,
                f"Toggler_{event.load}",
                1,
                f"Toggler_{event.load}",
                "PQ",
                event.load,
                event.time,
            ]
        )
        workbook.save(destination)
        return destination

    def _disable_togglers(self, workbook: Any) -> None:
        for sheet_name in ("Toggler", "Toggle"):
            if sheet_name not in workbook.sheetnames:
                continue
            sheet = workbook[sheet_name]
            headers = [cell.value for cell in sheet[1]]
            if "u" not in headers:
                continue
            u_col = headers.index("u") + 1
            for row in range(2, sheet.max_row + 1):
                sheet.cell(row=row, column=u_col, value=0)

    def _simulate(
        self,
        name: str,
        case_path: Path,
        final_time: float,
    ) -> SimulationSeries:
        import andes

        system = andes.load(str(case_path), setup=True, no_output=True)
        system.TDS.config.tf = final_time
        system.TDS.config.tstep = 1.0 / 30.0
        system.TDS.config.no_tqdm = 1
        system.PFlow.run()
        system.TDS.run(no_summary=True)
        frequency = system.TDS.get_timeseries(system.BusFreq.f)
        voltage = system.TDS.get_timeseries(system.Bus.v)
        gen_omega = system.TDS.get_timeseries(system.GENROU.omega)

        return SimulationSeries(
            name=name,
            frequency=frequency,
            voltage=voltage,
            gen_omega=gen_omega,
            regcp1_pe=self._optional_timeseries(system, "REGCP1", "Pe"),
            regf1_pe=self._optional_timeseries(system, "REGF1", "Pe"),
            final_time=float(frequency.index.max()) if not frequency.empty else 0.0,
        )

    def _optional_timeseries(
        self,
        system: Any,
        model_name: str,
        variable_name: str,
    ) -> pd.DataFrame | None:
        if not hasattr(system, model_name):
            return None
        model = getattr(system, model_name)
        if getattr(model, "n", 0) == 0 or not hasattr(model, variable_name):
            return None
        return system.TDS.get_timeseries(getattr(model, variable_name))

    def _write_outputs(
        self,
        output_dir: Path,
        base: SimulationSeries,
        ibr: SimulationSeries,
        profile: str,
        event: LoadToggleEvent,
    ) -> dict[str, Path]:
        artifacts = {
            "frequency_plot": output_dir / "frequency_busfreq.png",
            "voltage_plot": output_dir / "voltage_buses.png",
            "ibr_power_plot": output_dir / "ibr_power.png",
            "summary": output_dir / "summary.json",
            "base_frequency_csv": output_dir / "base_frequency.csv",
            "ibr_frequency_csv": output_dir / f"{profile}_frequency.csv",
            "base_voltage_csv": output_dir / "base_voltage.csv",
            "ibr_voltage_csv": output_dir / f"{profile}_voltage.csv",
        }

        base.frequency.to_csv(artifacts["base_frequency_csv"])
        ibr.frequency.to_csv(artifacts["ibr_frequency_csv"])
        base.voltage.to_csv(artifacts["base_voltage_csv"])
        ibr.voltage.to_csv(artifacts["ibr_voltage_csv"])

        self._plot_frequency(base, ibr, artifacts["frequency_plot"], event)
        self._plot_voltage(base, ibr, artifacts["voltage_plot"], event)
        self._plot_ibr_power(ibr, artifacts["ibr_power_plot"])
        artifacts["summary"].write_text(
            json.dumps(self._summary(base, ibr, event), indent=2),
            encoding="utf-8",
        )
        return artifacts

    def _plot_frequency(
        self,
        base: SimulationSeries,
        ibr: SimulationSeries,
        destination: Path,
        event: LoadToggleEvent,
    ) -> None:
        buses = ["BusFreq_1", "BusFreq_5", "BusFreq_6", "BusFreq_9", "BusFreq_10"]
        labels = {
            "BusFreq_1": "bus 30 SG",
            "BusFreq_5": "bus 34 GFL site",
            "BusFreq_6": "bus 35 GFM site",
            "BusFreq_9": "bus 38 GFL site",
            "BusFreq_10": "bus 39 ext",
        }
        fig, axes = plt.subplots(len(buses), 1, figsize=(9, 9), sharex=True)
        for axis, bus in zip(axes, buses, strict=True):
            if bus in base.frequency:
                axis.plot(
                    base.frequency.index,
                    base.frequency[bus],
                    label="base",
                    lw=1.2,
                )
            if bus in ibr.frequency:
                axis.plot(
                    ibr.frequency.index,
                    ibr.frequency[bus],
                    label=ibr.name,
                    lw=1.2,
                )
            axis.set_ylabel(labels[bus])
            axis.grid(True, alpha=0.25)
        axes[0].legend(loc="best")
        axes[-1].set_xlabel("time [s]")
        fig.suptitle(f"Bus frequency response after {event.label}")
        fig.tight_layout()
        fig.savefig(destination, dpi=180)
        plt.close(fig)

    def _plot_voltage(
        self,
        base: SimulationSeries,
        ibr: SimulationSeries,
        destination: Path,
        event: LoadToggleEvent,
    ) -> None:
        buses = [16, 30, 34, 35, 38, 39]
        fig, axes = plt.subplots(len(buses), 1, figsize=(9, 9), sharex=True)
        for axis, bus in zip(axes, buses, strict=True):
            if bus in base.voltage:
                axis.plot(base.voltage.index, base.voltage[bus], label="base", lw=1.2)
            if bus in ibr.voltage:
                axis.plot(ibr.voltage.index, ibr.voltage[bus], label=ibr.name, lw=1.2)
            axis.set_ylabel(f"V bus {bus}")
            axis.grid(True, alpha=0.25)
        axes[0].legend(loc="best")
        axes[-1].set_xlabel("time [s]")
        fig.suptitle(f"Bus voltage response after {event.label}")
        fig.tight_layout()
        fig.savefig(destination, dpi=180)
        plt.close(fig)

    def _plot_ibr_power(self, ibr: SimulationSeries, destination: Path) -> None:
        fig, axes = plt.subplots(2, 1, figsize=(9, 6), sharex=True)
        if ibr.regcp1_pe is not None:
            for column in ibr.regcp1_pe.columns:
                axes[0].plot(ibr.regcp1_pe.index, ibr.regcp1_pe[column], label=column)
        axes[0].set_ylabel("REGCP1 Pe")
        axes[0].grid(True, alpha=0.25)
        axes[0].legend(loc="best")
        if ibr.regf1_pe is not None:
            for column in ibr.regf1_pe.columns:
                axes[1].plot(ibr.regf1_pe.index, ibr.regf1_pe[column], label=column)
        axes[1].set_ylabel("REGF1 Pe")
        axes[1].set_xlabel("time [s]")
        axes[1].grid(True, alpha=0.25)
        axes[1].legend(loc="best")
        fig.suptitle(f"{ibr.name} IBR active power response")
        fig.tight_layout()
        fig.savefig(destination, dpi=180)
        plt.close(fig)

    def _summary(
        self,
        base: SimulationSeries,
        ibr: SimulationSeries,
        event: LoadToggleEvent,
    ) -> dict[str, Any]:
        return {
            "event": {
                "type": "toggler_load_disconnect",
                "model": "PQ",
                "load": event.load,
                "time": event.time,
                "label": event.label,
            },
            "base": {
                "final_time": base.final_time,
                "frequency_columns": list(base.frequency.columns),
                "voltage_columns": list(base.voltage.columns),
                "gen_omega_columns": list(base.gen_omega.columns),
            },
            "ibr": {
                "final_time": ibr.final_time,
                "frequency_columns": list(ibr.frequency.columns),
                "voltage_columns": list(ibr.voltage.columns),
                "gen_omega_columns": list(ibr.gen_omega.columns),
                "regcp1_columns": []
                if ibr.regcp1_pe is None
                else list(ibr.regcp1_pe.columns),
                "regf1_columns": []
                if ibr.regf1_pe is None
                else list(ibr.regf1_pe.columns),
            },
            "completed": self._completed(base, ibr, event.final_time),
            "metrics": self._metrics(base, ibr),
        }

    def _metrics(
        self,
        base: SimulationSeries,
        ibr: SimulationSeries,
    ) -> dict[str, float]:
        metrics: dict[str, float] = {
            "base_final_time": base.final_time,
            f"{ibr.name}_final_time": ibr.final_time,
        }
        for bus in ("BusFreq_1", "BusFreq_5", "BusFreq_6", "BusFreq_9", "BusFreq_10"):
            if bus in base.frequency:
                metrics[f"base_{bus}_min_frequency"] = float(base.frequency[bus].min())
                metrics[f"base_{bus}_max_abs_frequency_delta"] = float(
                    (base.frequency[bus] - base.frequency[bus].iloc[0]).abs().max()
                )
            if bus in ibr.frequency:
                metrics[f"{ibr.name}_{bus}_min_frequency"] = float(
                    ibr.frequency[bus].min()
                )
                metrics[f"{ibr.name}_{bus}_max_abs_frequency_delta"] = float(
                    (ibr.frequency[bus] - ibr.frequency[bus].iloc[0]).abs().max()
                )
        for bus in (16, 30, 34, 35, 38, 39):
            if bus in base.voltage:
                metrics[f"base_bus_{bus}_max_abs_voltage_delta"] = float(
                    (base.voltage[bus] - base.voltage[bus].iloc[0]).abs().max()
                )
            if bus in ibr.voltage:
                metrics[f"{ibr.name}_bus_{bus}_max_abs_voltage_delta"] = float(
                    (ibr.voltage[bus] - ibr.voltage[bus].iloc[0]).abs().max()
                )
        if ibr.regcp1_pe is not None:
            metrics[f"{ibr.name}_regcp1_max_abs_pe_delta"] = float(
                (ibr.regcp1_pe - ibr.regcp1_pe.iloc[0]).abs().max().max()
            )
        if ibr.regf1_pe is not None:
            metrics[f"{ibr.name}_regf1_max_abs_pe_delta"] = float(
                (ibr.regf1_pe - ibr.regf1_pe.iloc[0]).abs().max().max()
            )
        return metrics

    def _completed(
        self,
        base: SimulationSeries,
        ibr: SimulationSeries,
        expected_final_time: float,
    ) -> bool:
        tolerance = 2.0 / 30.0
        return (
            base.final_time >= expected_final_time - tolerance
            and ibr.final_time >= expected_final_time - tolerance
        )

    def _resolve(self, path: Path) -> Path:
        if path.is_absolute():
            return path
        return self.project_root / path
